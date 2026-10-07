"""Извлечение текста из загруженных файлов и нарезка на фрагменты."""
import html
import io
import re
import zipfile
from html.parser import HTMLParser

from . import config


class IngestError(Exception):
    pass


def _decode(data: bytes) -> str:
    if b"\x00" in data[:4096]:
        # возможно UTF-16
        for enc in ("utf-16", "utf-16-le"):
            try:
                t = data.decode(enc)
                if t.strip():
                    return t
            except Exception:
                pass
        raise IngestError("Файл бинарный, а не текстовый")
    for enc in ("utf-8-sig", "cp1251", "cp866", "latin-1"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


class _Strip(HTMLParser):
    BLOCK = {"p", "div", "br", "li", "tr", "h1", "h2", "h3", "h4", "h5", "h6", "section", "article", "pre"}

    def __init__(self):
        super().__init__()
        self.out, self.skip = [], 0

    def handle_starttag(self, tag, attrs):
        if tag in ("script", "style", "noscript", "svg"):
            self.skip += 1
        if tag in self.BLOCK:
            self.out.append("\n")

    def handle_endtag(self, tag):
        if tag in ("script", "style", "noscript", "svg"):
            self.skip = max(0, self.skip - 1)
        if tag in self.BLOCK:
            self.out.append("\n")

    def handle_data(self, data):
        if not self.skip:
            self.out.append(data)


def _html_to_text(raw: str) -> str:
    p = _Strip()
    p.feed(raw)
    return "".join(p.out)


def _pdf(data: bytes) -> str:
    text = ""
    try:
        import pymupdf  # type: ignore
        doc = pymupdf.open(stream=data, filetype="pdf")
        text = "\n\n".join(page.get_text() for page in doc)
    except Exception:
        text = ""
    if len(text.strip()) < 40:
        try:
            import pypdf
            r = pypdf.PdfReader(io.BytesIO(data))
            text = "\n\n".join((p.extract_text() or "") for p in r.pages)
        except Exception as e:
            raise IngestError(f"Не удалось прочитать PDF: {e}")
    if len(text.strip()) < 40:
        raise IngestError("В PDF нет текстового слоя (скан?). Нужен распознанный текст.")
    return text


def _docx(data: bytes) -> str:
    import docx
    d = docx.Document(io.BytesIO(data))
    parts = []
    for p in d.paragraphs:
        t = p.text.strip()
        if not t:
            continue
        style = (p.style.name or "").lower() if p.style is not None else ""
        if style.startswith("heading") or style.startswith("заголовок") or style == "title":
            parts.append("\n## " + t)
        else:
            parts.append(t)
    for tb in d.tables:
        for row in tb.rows:
            parts.append(" | ".join(c.text.strip() for c in row.cells))
    return "\n".join(parts)


def _pptx(data: bytes) -> str:
    z = zipfile.ZipFile(io.BytesIO(data))
    slides = sorted((n for n in z.namelist() if re.match(r"ppt/slides/slide\d+\.xml$", n)),
                    key=lambda n: int(re.findall(r"\d+", n)[-1]))
    out = []
    for i, n in enumerate(slides, 1):
        xml = z.read(n).decode("utf-8", "replace")
        texts = re.findall(r"<a:t>(.*?)</a:t>", xml, flags=re.S)
        out.append(f"## Слайд {i}\n" + "\n".join(html.unescape(t) for t in texts))
    return "\n\n".join(out)


def _odt(data: bytes) -> str:
    z = zipfile.ZipFile(io.BytesIO(data))
    xml = z.read("content.xml").decode("utf-8", "replace")
    xml = re.sub(r"</text:(p|h)>", "\n", xml)
    return html.unescape(re.sub(r"<[^>]+>", "", xml))


def _rtf(raw: str) -> str:
    raw = re.sub(r"\\'([0-9a-fA-F]{2})", lambda m: bytes([int(m.group(1), 16)]).decode("cp1251", "replace"), raw)
    raw = re.sub(r"\\u(-?\d+)\??", lambda m: chr(int(m.group(1)) % 65536), raw)
    raw = re.sub(r"\\par[d]?", "\n", raw)
    raw = re.sub(r"\\[a-zA-Z]+-?\d* ?", "", raw)
    return re.sub(r"[{}]", "", raw)


def extract_text(filename: str, data: bytes) -> str:
    ext = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""
    if ext == "pdf" or data[:5] == b"%PDF-":
        text = _pdf(data)
    elif ext == "docx":
        text = _docx(data)
    elif ext == "pptx":
        text = _pptx(data)
    elif ext in ("odt", "odp"):
        text = _odt(data)
    elif ext in ("doc", "xls", "xlsx", "ppt", "zip", "rar", "7z", "png", "jpg", "jpeg", "gif", "exe", "mp3", "mp4"):
        raise IngestError(f"Формат .{ext} не поддерживается (нужен текст, PDF или DOCX)")
    else:
        raw = _decode(data)
        if ext in ("html", "htm", "xhtml"):
            text = _html_to_text(raw)
        elif ext == "rtf" or raw.startswith("{\\rtf"):
            text = _rtf(raw)
        elif raw.lstrip().lower().startswith(("<!doctype html", "<html")):
            text = _html_to_text(raw)
        else:
            text = raw
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\x00", "")
    text = re.sub(r"[ \t\u00a0]+", " ", text)
    text = re.sub(r"-\n(?=[a-zа-яё])", "", text)  # переносы слов в PDF
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    if len(text) < 20:
        raise IngestError("В файле почти нет текста")
    return text


def chunk_text(text: str, size: int = config.CHUNK_CHARS) -> list[str]:
    """Нарезка по абзацам примерно до size символов."""
    paras = [p.strip() for p in re.split(r"\n\s*\n", text) if p.strip()]
    out, cur = [], ""
    for p in paras:
        while len(p) > size * 1.5:  # очень длинный абзац режем по предложениям/символам
            cut = p.rfind(". ", 0, size)
            cut = cut + 1 if cut > size // 2 else size
            if cur:
                out.append(cur)
                cur = ""
            out.append(p[:cut].strip())
            p = p[cut:].strip()
        if len(cur) + len(p) + 2 > size and cur:
            out.append(cur)
            cur = p
        else:
            cur = (cur + "\n\n" + p) if cur else p
    if cur:
        out.append(cur)
    return out


def detect_language(text: str) -> str:
    sample = text[:6000]
    cyr = len(re.findall(r"[а-яёА-ЯЁ]", sample))
    lat = len(re.findall(r"[a-zA-Z]", sample))
    return "ru" if cyr >= lat * 0.5 else "en"
